import csv
import io
from uuid import uuid4

from conftest import TEST_PASSWORD
from openpyxl import load_workbook

from app.models import Candidate
from app.schemas import CandidateCreate
from app.services.deduplication import deduplicate_candidates

BASE = "/api/v1"
BRIEF = {
    "title": "Vision experts",
    "field": "Artificial Intelligence",
    "location": "Istanbul",
    "target_type": "Speaker",
    "requested_quantity": 20,
    "keywords": ["Computer Vision"],
}
PERSON = {
    "full_name": "Example Person",
    "organization": "Example University",
    "job_title": "Professor",
    "field": "Artificial Intelligence",
    "location": "Istanbul",
    "bio": "Computer Vision researcher",
    "expertise": ["Computer Vision"],
    "source_url": "https://example.com/profile",
    "source_name": "Example University",
}


def request(client):
    response = client.post(BASE + "/research-requests", json=BRIEF)
    assert response.status_code == 201, response.text
    return response.json()["id"]


def add(client, rid, person=None):
    response = client.post(f"{BASE}/research-requests/{rid}/candidates", json=person or PERSON)
    assert response.status_code == 201, response.text
    return response.json()["candidate"]


def test_complete_review_export_flow(auth):
    rid = request(auth)
    c = add(auth, rid)
    assert c["status"] == "NEW" and c["sources"][0]["url"] == PERSON["source_url"]
    assert c["relevance_score"] == 100
    assert c["confidence_score"] > 0
    before = auth.post(f"{BASE}/research-requests/{rid}/exports", json={"format": "json"}).json()
    assert before == []
    response = auth.patch(
        f"{BASE}/candidates/{c['id']}",
        json={"status": "APPROVED", "notes": "Reviewed against source"},
    )
    assert response.status_code == 200
    assert response.json()["status"] == "APPROVED"
    for format in ["csv", "xlsx", "json"]:
        exported = auth.post(f"{BASE}/research-requests/{rid}/exports", json={"format": format})
        assert exported.status_code == 200
        if format == "json":
            assert exported.json()[0]["Notes"] == "Reviewed against source"
        elif format == "csv":
            assert (
                list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))[0]["Name"]
                == PERSON["full_name"]
            )
        else:
            assert load_workbook(io.BytesIO(exported.content)).active.max_row == 2
    assert auth.get(BASE + "/dashboard").json()["approved"] == 1
    assert auth.get(BASE + "/exports").json()["total"] == 4


def test_n8n_import_completion_and_authentication(service):
    rid = request(service)
    c = add(service, rid)
    assert (
        service.patch(f"{BASE}/candidates/{c['id']}", json={"status": "APPROVED"}).status_code
        == 403
    )
    assert (
        service.patch(
            f"{BASE}/candidates/{c['id']}", json={"notes": "Additional evidence"}
        ).status_code
        == 200
    )
    assert (
        service.post(f"{BASE}/research-requests/{rid}/complete").json()["status"]
        == "READY_FOR_REVIEW"
    )
    assert (
        service.get(f"{BASE}/research-requests/{rid}/candidates").json()["items"][0]["status"]
        == "NEW"
    )
    assert service.get(BASE + "/admin/users").status_code == 403
    assert (
        service.post(f"{BASE}/research-requests/{rid}/exports", json={"format": "csv"}).status_code
        == 403
    )


def test_bulk_dedup_retains_sources_and_review(auth):
    rid = request(auth)
    c = add(auth, rid)
    auth.patch(
        f"{BASE}/candidates/{c['id']}", json={"status": "APPROVED", "notes": "Keep this review"}
    )
    payload = {
        "candidates": [
            {**PERSON, "source_url": "https://example.org/person"},
            {**PERSON, "full_name": "Different Person", "source_url": "https://example.com/other"},
        ]
    }
    result = auth.post(f"{BASE}/research-requests/{rid}/candidates/bulk", json=payload)
    assert result.status_code == 200, result.text
    assert result.json()["created"] == 1 and result.json()["merged"] == 1
    again = auth.post(f"{BASE}/research-requests/{rid}/candidates/bulk", json=payload).json()
    assert again["created"] == 0 and again["merged"] == 2
    saved = auth.get(f"{BASE}/candidates/{c['id']}").json()
    assert (
        saved["status"] == "APPROVED"
        and saved["notes"] == "Keep this review"
        and len(saved["sources"]) == 2
    )
    other = request(auth)
    assert add(auth, other)["id"] != c["id"]


def test_invalid_payloads_and_injection(auth):
    rid = request(auth)
    for payload in [
        {**PERSON, "full_name": ""},
        {**PERSON, "source_url": "javascript:alert(1)"},
        {**PERSON, "confidence_score": 101},
        {**PERSON, "status": "APPROVED"},
        {"full_name": "No source"},
        {**PERSON, "website": "https://user:secret@example.com"},
    ]:
        response = auth.post(f"{BASE}/research-requests/{rid}/candidates", json=payload)
        assert response.status_code == 422, response.text
    assert (
        auth.post(
            f"{BASE}/research-requests/{rid}/candidates/bulk",
            json={"candidates": [PERSON, {"full_name": "Invalid"}]},
        ).status_code
        == 422
    )
    assert auth.get(f"{BASE}/research-requests/{rid}/candidates").json()["total"] == 0
    c = add(auth, rid, {**PERSON, "full_name": '=HYPERLINK("bad")'})
    assert auth.patch(f"{BASE}/candidates/{c['id']}", json={"notes": None}).status_code == 422
    exported = auth.post(
        f"{BASE}/research-requests/{rid}/exports", json={"format": "csv", "include_all": True}
    )
    assert list(csv.DictReader(io.StringIO(exported.content.decode("utf-8-sig"))))[0][
        "Name"
    ].startswith("'=")
    assert (
        auth.get(BASE + "/candidates", params={"q": "'; DROP TABLE candidates;--"}).json()["total"]
        == 0
    )


def test_login_logout_role_and_csrf(client):
    assert client.get(BASE + "/research-requests").status_code == 401
    assert (
        client.get(
            BASE + "/research-requests", headers={"Authorization": "Bearer wrong"}
        ).status_code
        == 401
    )
    assert (
        client.post(
            BASE + "/auth/login", json={"email": "reviewer@example.com", "password": TEST_PASSWORD}
        ).status_code
        == 403
    )
    client.headers["X-Requested-With"] = "RelationshipPlatform"
    assert (
        client.post(
            BASE + "/auth/login", json={"email": "reviewer@example.com", "password": TEST_PASSWORD}
        ).status_code
        == 200
    )
    assert client.get(BASE + "/admin/users").status_code == 403
    cookie = client.cookies.get("rip_session")
    assert client.post(BASE + "/auth/logout").status_code == 200
    client.cookies.set("rip_session", cookie)
    assert client.get(BASE + "/auth/me").status_code == 401


def test_demo_run_idempotent_and_unverified(auth):
    rid = request(auth)
    result = auth.post(f"{BASE}/research-requests/{rid}/run")
    assert result.status_code == 200, result.text
    assert result.json()["created"] == 12
    assert auth.post(f"{BASE}/research-requests/{rid}/run").json()["created"] == 0
    candidates = auth.get(f"{BASE}/research-requests/{rid}/candidates").json()["items"]
    assert all(
        c["is_demo"] and c["confidence_score"] == 0 and c["status"] == "NEW" for c in candidates
    )
    assert auth.get(f"{BASE}/research-requests/{rid}").json()["status"] == "READY_FOR_REVIEW"


def test_pagination_filters_sort_and_not_found(auth):
    rid = request(auth)
    auth.post(f"{BASE}/research-requests/{rid}/run")
    page = auth.get(f"{BASE}/candidates?page_size=5&page=2&sort=full_name&direction=asc").json()
    assert len(page["items"]) == 5 and page["total"] == 12
    c = page["items"][0]
    auth.patch(f"{BASE}/candidates/{c['id']}", json={"status": "REJECTED"})
    assert auth.get(BASE + "/candidates?status=REJECTED").json()["total"] == 1
    assert auth.get(BASE + "/candidates?awaiting_review=true").json()["total"] == 11
    assert auth.get(BASE + "/candidates?q=NotAName").json()["total"] == 0
    assert auth.get(BASE + "/candidates?min_score=101").status_code == 422
    assert auth.get(BASE + "/candidates?sort=password_hash").status_code == 422
    assert auth.get(f"{BASE}/research-requests/{uuid4()}").status_code == 404
    assert auth.get(BASE + "/research-requests/not-a-uuid").status_code == 422


def test_matching_requires_identity_evidence():
    existing = Candidate(
        full_name="Alice Example", organization="Shared Company", website="https://example.com"
    )
    unrelated = CandidateCreate(
        full_name="Bob Else",
        organization="Shared Company",
        website="https://example.com",
        source_url="https://example.com/bob",
    )
    assert deduplicate_candidates([existing], unrelated) is None
    variant = CandidateCreate(
        full_name="  ALICE EXAMPLE  ",
        organization="Shared Company",
        source_url="https://example.com/alice",
    )
    assert deduplicate_candidates([existing], variant) is existing


def test_source_edit_and_research_update(auth):
    rid = request(auth)
    c = add(auth, rid)
    response = auth.patch(
        f"{BASE}/candidates/{c['id']}",
        json={
            "sources": [
                {"url": PERSON["source_url"], "name": "Updated source", "source_type": "university"}
            ]
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["sources"][0]["name"] == "Updated source"
    assert auth.patch(f"{BASE}/candidates/{c['id']}", json={"sources": []}).status_code == 422
    result = auth.patch(
        f"{BASE}/research-requests/{rid}",
        json={**BRIEF, "location": "Berlin", "status": "COMPLETED"},
    )
    assert result.status_code == 200
    # Stated in Istanbul for a Berlin brief: the 20 place points are lost and a conflicting
    # place costs 20 more, so it ranks below a person whose place is unknown.
    assert auth.get(f"{BASE}/candidates/{c['id']}").json()["relevance_score"] == 60
    assert auth.delete(f"{BASE}/research-requests/{rid}").status_code == 204
    assert auth.get(f"{BASE}/candidates/{c['id']}").status_code == 404


def test_mock_integrations_and_user_management(auth):
    rid = request(auth)
    response = auth.post(
        f"{BASE}/research-requests/{rid}/exports", json={"format": "sheets"}
    ).json()
    assert response["provider"] == "mock" and response["url"] is None
    assert auth.post(BASE + "/ai/interpret", json=BRIEF).json()["keywords"] == [
        "Artificial Intelligence"
    ]
    user = auth.post(
        BASE + "/admin/users",
        json={"email": "new@example.com", "name": "New User", "password": TEST_PASSWORD},
    ).json()
    assert "password_hash" not in user
    assert auth.post(f"{BASE}/admin/users/{user['id']}/toggle").json()["active"] is False
    assert auth.get(BASE + "/admin/audit").status_code == 200
